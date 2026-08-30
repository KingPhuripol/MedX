#!/usr/bin/env python3
"""สร้างไฟล์ .docx ของเอกสาร Senior Project IDEA จากต้นฉบับ markdown

ฟอร์แมตยึดตาม IdeaDocument_Regular_Example2.pdf ที่คณะให้มา — TH SarabunPSK 16pt

ข้อกำหนดของต้นฉบับ: แต่ละย่อหน้าต้องอยู่บรรทัดเดียว ห้ามตัดบรรทัดกลางย่อหน้า
เพราะภาษาไทยไม่เว้นวรรคระหว่างคำ การตัดบรรทัดจะกลายเป็นช่องว่างแปลกปลอมในไฟล์ที่ได้
(ตัวเชื่อมบรรทัดในไฟล์นี้ป้องกันไว้แล้ว แต่ต้นฉบับบรรทัดเดียวอ่านง่ายกว่า)

ต้องพึ่ง python-docx ซึ่งเป็นไลบรารีภายนอก จึงไม่อยู่ใน scripts/ ที่เป็น stdlib เท่านั้น
ติดตั้ง: pip install python-docx

ใช้งาน: python3 tools/build_idea_docx.py <ต้นฉบับ.md> <ผลลัพธ์.docx>
"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn

SRC = sys.argv[1]
OUT = sys.argv[2]
FONT = "TH SarabunPSK"
SIZE = 16

# ---------- parse markdown ----------
raw = open(SRC, encoding="utf-8").read()
lines = [l for l in raw.split("\n") if not l.startswith(">")]
raw = "\n".join(lines)
raw = re.sub(r"^# .*\n", "", raw)                      # drop md title
body, _, sig = raw.partition("\n---\n")

sections = {}
order = []
for m in re.finditer(r"^## (.+?)\n(.*?)(?=^## |\Z)", body, re.S | re.M):
    sections[m.group(1).strip()] = m.group(2).strip()
    order.append(m.group(1).strip())

THAI = re.compile(r"[฀-๿]")

def join_lines(lines):
    """เชื่อมบรรทัดแบบรู้ภาษาไทย — ไทยชนไทยไม่ใส่ช่องว่าง เพราะไทยไม่เว้นวรรคระหว่างคำ"""
    out = ""
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        if out and not (THAI.match(out[-1]) and THAI.match(ln[0])):
            out += " "
        out += ln
    return out

def blocks(text):
    """แยกเป็นบล็อกตามบรรทัดว่าง แล้วรวมบรรทัดในบล็อกเดียวกัน"""
    return [join_lines(b.split("\n")) for b in re.split(r"\n\s*\n", text) if b.strip()]

def numbered(text):
    """แยกรายการเลข รองรับบรรทัดต่อเนื่องที่ย่อหน้าเข้า"""
    items, cur = [], None
    for line in text.split("\n"):
        if re.match(r"^\d+\.\s", line.strip()):
            if cur: items.append(cur)
            cur = re.sub(r"^\d+\.\s*", "", line.strip())
        elif line.strip() and cur is not None:
            nxt = line.strip()
            sep = "" if (THAI.match(cur[-1]) and THAI.match(nxt[0])) else " "
            cur += sep + nxt
    if cur: items.append(cur)
    return items

# ---------- docx helpers ----------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.54))

def style_run(run, bold=False, underline=False, size=SIZE):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    rPr = run._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT)
    rPr.append(rPr.makeelement(qn("w:szCs"), {qn("w:val"): str(size * 2)}))
    rPr.append(rPr.makeelement(qn("w:lang"),
                               {qn("w:val"): "en-US", qn("w:bidi"): "th-TH"}))
    if bold:
        rPr.append(rPr.makeelement(qn("w:bCs"), {}))

INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")


def add_rich(p, text, size=SIZE):
    """เติมข้อความลงย่อหน้า โดยแปลง **ตัวหนา** และ *ตัวเอียง* เป็นรูปแบบจริง"""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            style_run(p.add_run(text[pos:m.start()]), size=size)
        if m.group(1) is not None:
            style_run(p.add_run(m.group(1)), bold=True, size=size)
        else:
            run = p.add_run(m.group(2))
            style_run(run, size=size)
            run.italic = True
        pos = m.end()
    if pos < len(text):
        style_run(p.add_run(text[pos:]), size=size)


def para(text="", *, bold=False, underline=False, indent=0.0, first=0.0,
         align=None, space_after=0, size=SIZE):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(first)
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    if text:
        if "*" in text and not bold and not underline:
            add_rich(p, text, size=size)
        else:
            style_run(p.add_run(text), bold=bold, underline=underline, size=size)
    return p

def heading(text):
    para(text, underline=True, space_after=0)

# ---------- build ----------
para("Senior Project IDEA", underline=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

for name in order:
    text = sections[name]
    if name == "References":
        # ให้ References ขึ้นหน้าใหม่ เพื่อให้เห็นชัดว่าเนื้อหาจบที่หน้าใด
        # เกณฑ์ของคณะคือ "2-4 หน้า บวกเอกสารอ้างอิง" — อ้างอิงไม่นับรวมในโควตาหน้า
        doc.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)
    heading(name)

    if name in ("Topic", "Advisor"):
        for b in blocks(text):
            para(b, indent=1.27, space_after=0)

    elif name == "Members":
        for i, item in enumerate(numbered(text), 1):
            para(f"{i}. {item}", indent=1.9, space_after=0)

    elif name == "Objective":
        for i, item in enumerate(numbered(text), 1):
            p = para(f"{i}. {item}", indent=1.9, space_after=0,
                     align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.first_line_indent = Cm(-0.75)

    elif name.startswith("Idea of Project"):
        for b in blocks(text):
            para(b, first=1.27, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0)

    elif name == "References":
        for b in blocks(text):
            ref = para(space_after=4)
            ref.paragraph_format.left_indent = Cm(1.27)
            ref.paragraph_format.first_line_indent = Cm(-1.27)
            add_rich(ref, b)

    else:  # Scope of work / Expected Result
        for b in blocks(text):
            m = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", b, re.S)
            if m:
                para(m.group(1).strip(), bold=True, space_after=0)
                if m.group(2).strip():
                    para(m.group(2).strip(), align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         space_after=4)
            else:
                para(b, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4)

    doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

# ---------- signature ----------
para(space_after=0)
for line in [l.strip() for l in sig.split("\n") if l.strip()]:
    para(line, align=WD_ALIGN_PARAGRAPH.CENTER, indent=6.0, space_after=0)

doc.save(OUT)
print("saved:", OUT)
